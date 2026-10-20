# -*- coding: utf-8 -*-
"""
entregas_materiales.py - registro de entregas de material de una obra

ESPECIFICACION (piloto Olabeaga, 20/10/2026; aprobada por Bixente: "opcion A").

Problema que resuelve
  El Excel de materiales era una rejilla (catalogo entero x una columna por
  fecha) sin sitio para suministrador ni albaran, y el panel solo leia su
  ultima hoja. Dos formatos para lo mismo, y una rejilla de 117 filas con
  formulas que se rompio al editarla desde codigo.

Diseno
  * UNA hoja ``ENTREGAS`` por obra, con una fila por linea entregada:
        FECHA | SUMINISTRADOR | ALBARAN | CATEGORIA | MATERIAL | TIPO | MARCA |
        UNI | CANTIDAD | NOTAS
    Es lo UNICO que se escribe (a mano o por agy). Todo lo demas se genera.
  * El panel lee ENTREGAS directamente (``leer_materiales_wb``), asi que
    siempre esta al dia. El Excel lleva ademas una hoja por mes y un
    ACUMULADO, generadas de la misma agrupacion (``construir_hojas``): lo que
    ves en pantalla y en el Excel es el mismo calculo.
  * Cada mes lista SOLO lo entregado ese mes (no repite el catalogo): una LISTA
    con una fila por linea entregada (fecha, suministrador y albaran como
    COLUMNAS, mas material, tipo, uni y cantidad) y debajo el TOTAL DEL MES POR
    MATERIAL (MES ANT, TOTAL MES, ACUMULADO). Las cabeceras se pulsan para
    ordenar en el panel y llevan filtros de Excel en la hoja.
  * El "Informe de obra" a la carta ofrece Total acumulado, Por mes, Por
    albaran, Por suministrador y Todas las entregas (secciones_informe).
  * Las hojas generadas llevan solo valores, sin formulas: no existe el
    "total sin calcular" que dejo a Olabeaga fuera del portal.

Reglas de datos (CLAUDE.md: no inventar)
  * Una fila invalida (fecha ilegible, cantidad no numerica o ambigua) queda
    FUERA y se avisa con su numero de fila; nunca se convierte en cero.
  * Suministrador o albaran desconocidos son ``None`` ("sin dato"), no "0".
  * Dos filas identicas se cuentan las dos pero se avisan (entrada doble).
"""
from __future__ import annotations

import html as html_lib
import io
import os
import re
import shutil
import unicodedata
from datetime import date, datetime

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.datetime import from_excel
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # el resto del motor ya avisa si falta openpyxl
    openpyxl = None

HOJA_ENTREGAS = 'ENTREGAS'
HOJA_ACUMULADO = 'ACUMULADO'
HOJA_LEEME = 'LEEME'
CABECERA = ['FECHA', 'SUMINISTRADOR', 'ALBARÁN', 'CATEGORÍA', 'MATERIAL',
            'TIPO', 'MARCA', 'UNI', 'CANTIDAD', 'NOTAS']
CAMPOS = ['fecha', 'suministrador', 'albaran', 'categoria', 'material',
          'tipo', 'marca', 'uni', 'cantidad', 'notas']
OBLIGATORIOS = ('fecha', 'material', 'cantidad')
MESES = ['Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio', 'Julio',
         'Agosto', 'Septiembre', 'Octubre', 'Noviembre', 'Diciembre']
_RE_HOJA_MES = re.compile(r'^(%s)\d{2}$' % '|'.join(MESES))
COMANDO_GENERAR = 'py -3.11 _SISTEMA/MOTOR/scripts/materiales_xlsx.py generar <obra>'
MAX_AVISOS_EN_PANEL = 5


# --------------------------------------------------------------- utilidades

def _sin_tildes(s: str) -> str:
    return ''.join(c for c in unicodedata.normalize('NFD', s)
                   if unicodedata.category(c) != 'Mn')


def _norm_cabecera(v) -> str:
    return re.sub(r'[^A-Z0-9 ]', ' ', _sin_tildes(str(v or '')).upper()).strip()


def _campo_de_cabecera(v):
    n = _norm_cabecera(v)
    if not n:
        return None
    if n.startswith('FECHA'):
        return 'fecha'
    if 'SUMINIST' in n or 'PROVEED' in n:
        return 'suministrador'
    if 'ALBARAN' in n:
        return 'albaran'
    if n.startswith('CATEGORIA'):
        return 'categoria'
    if n == 'MATERIAL':
        return 'material'
    if n == 'TIPO':
        return 'tipo'
    if n == 'MARCA':
        return 'marca'
    if n in ('UNI', 'UNIDAD', 'UNIDADES', 'UD', 'UDS'):
        return 'uni'
    if n.startswith('CANT'):
        return 'cantidad'
    if n.startswith('NOTA') or n.startswith('OBSERV'):
        return 'notas'
    return None


def _txt(v):
    """Texto limpio o None. Los numeros salen sin '.0' y con coma decimal."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        if float(v).is_integer():
            return str(int(v))
        return str(v).replace('.', ',')
    s = ' '.join(str(v).split())
    return s or None


def _num(v):
    """int/float limpio. Entero si no tiene decimales; sin ruido de coma flotante."""
    v = round(float(v), 6)
    return int(v) if v.is_integer() else v


def _fecha(v):
    """date o None. Acepta fecha de Excel, texto DD/MM/AAAA, AAAA-MM-DD y serie."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool) and 30000 < v < 80000:
        try:
            return from_excel(v).date()
        except (ValueError, OverflowError):
            return None
    if isinstance(v, str):
        s = v.strip()
        m = re.fullmatch(r'(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4})', s)
        if m:
            d, mo, a = map(int, m.groups())
        else:
            m = re.fullmatch(r'(\d{4})-(\d{1,2})-(\d{1,2})', s)
            if not m:
                return None
            a, mo, d = map(int, m.groups())
        try:
            return date(a, mo, d)
        except ValueError:
            return None
    return None


def _cantidad(v):
    """(numero, motivo_de_rechazo). Texto numerico simple vale; '1.200' es
    ambiguo (¿1,2 o 1200?) y se rechaza en vez de adivinar."""
    if isinstance(v, bool) or v is None:
        return None, 'esta vacia'
    if isinstance(v, (int, float)):
        return _num(v), None
    s = str(v).strip()
    if re.fullmatch(r'-?\d+[.,]\d{3}', s):
        return None, f"es ambigua ('{s}': ¿decimales o miles?)"
    if re.fullmatch(r'-?\d+([.,]\d+)?', s):
        return _num(s.replace(',', '.')), None
    return None, f"no es un numero ('{s}')"


def _clave_texto(v) -> str:
    return ' '.join(str(v or '').split()).upper()


def _fmt(n) -> str:
    if n is None:
        return ''
    if float(n).is_integer():
        return f"{int(n):,}".replace(',', '.')
    t = f"{n:,.2f}".rstrip('0').rstrip('.')
    return t.replace(',', '§').replace('.', ',').replace('§', '.')


SUMINISTRADORES_HABITUALES = ['Sagarde', 'Electrovizcaya', 'Suelbat']
_EN_MANO = ('SAGARDE', 'EN MANO', 'MANO')


def entrega_completa(suministrador, albaran) -> bool:
    """Regla de Bixente (20/10/2026): entrega en mano = suministrador Sagarde (no
    hay albaran); si no, la empresa y el codigo que vienen en el albaran. Una
    entrega con el dato incompleto se ve y se avisa; nunca se rellena sola."""
    s = _clave_texto(suministrador)
    if not s:
        return False
    return s in _EN_MANO or bool(_clave_texto(albaran))


def contar_incompletas(ruta) -> int:
    """Lineas de ENTREGAS sin suministrador o sin albaran (segun la regla)."""
    ent, _ = leer_registro(_abrir(ruta, data_only=True))
    return sum(1 for e in ent if not entrega_completa(e['suministrador'], e['albaran']))


def etiqueta_mes(f: date) -> str:
    return f"{MESES[f.month - 1]}{f.year % 100:02d}"


def _e(v) -> str:
    return html_lib.escape(str(v))


# ------------------------------------------------------------------- lectura

def leer_registro(wb, hoy: date | None = None):
    """Lee la hoja ENTREGAS de un libro de openpyxl.

    Devuelve ``(entregas, avisos)``. Cada entrega es un dict con ``fila`` (la
    del Excel), ``fecha`` (date), ``suministrador``, ``albaran``, ``categoria``,
    ``material``, ``tipo``, ``marca``, ``uni``, ``cantidad`` y ``notas``.
    """
    hoy = hoy or date.today()
    if HOJA_ENTREGAS not in wb.sheetnames:
        return [], [f"No existe la hoja {HOJA_ENTREGAS}."]
    ws = wb[HOJA_ENTREGAS]
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        return [], [f"La hoja {HOJA_ENTREGAS} esta vacia (falta la cabecera)."]

    columna = {}
    for i, h in enumerate(filas[0]):
        c = _campo_de_cabecera(h)
        if c and c not in columna:
            columna[c] = i
    faltan = [c.upper() for c in OBLIGATORIOS if c not in columna]
    if faltan:
        return [], [f"En {HOJA_ENTREGAS} falta la columna {', '.join(faltan)} en la fila 1."]

    entregas, avisos, vistas = [], [], {}
    for n, valores in enumerate(filas[1:], start=2):
        def v(campo):
            i = columna.get(campo)
            return valores[i] if i is not None and i < len(valores) else None
        if all(x in (None, '') for x in valores):
            continue
        f = _fecha(v('fecha'))
        if f is None:
            avisos.append(f"{HOJA_ENTREGAS} fila {n}: la fecha no es valida ({v('fecha')!r}); la fila no se cuenta.")
            continue
        material = _txt(v('material'))
        if not material:
            avisos.append(f"{HOJA_ENTREGAS} fila {n}: falta el material; la fila no se cuenta.")
            continue
        cant, motivo = _cantidad(v('cantidad'))
        if motivo:
            avisos.append(f"{HOJA_ENTREGAS} fila {n}: la cantidad {motivo}; la fila no se cuenta.")
            continue
        if cant == 0:
            avisos.append(f"{HOJA_ENTREGAS} fila {n}: cantidad 0; se ignora.")
            continue
        if f > hoy:
            avisos.append(f"{HOJA_ENTREGAS} fila {n}: fecha futura ({f:%d/%m/%Y}); se cuenta igualmente.")
        e = {'fila': n, 'fecha': f, 'suministrador': _txt(v('suministrador')),
             'albaran': _txt(v('albaran')), 'categoria': _txt(v('categoria')),
             'material': material, 'tipo': _txt(v('tipo')), 'marca': _txt(v('marca')),
             'uni': _txt(v('uni')), 'cantidad': cant, 'notas': _txt(v('notas'))}
        firma = (f, _clave_texto(e['albaran']), _clave_texto(e['suministrador']),
                 _clave_texto(e['categoria']), _clave_texto(material), _clave_texto(e['tipo']),
                 _clave_texto(e['marca']), _clave_texto(e['uni']), cant)
        if firma in vistas:
            avisos.append(f"{HOJA_ENTREGAS} filas {vistas[firma]} y {n}: linea repetida (misma fecha, albaran, "
                          "material y cantidad); se cuentan las dos, revisa si es una entrada doble.")
        else:
            vistas[firma] = n
        entregas.append(e)
    return entregas, avisos


# ---------------------------------------------------------------- agrupacion

def agrupar(entregas):
    """Agrupa las entregas por mes. Ver la ESPECIFICACION del modulo."""
    ent = sorted(entregas, key=lambda e: (e['fecha'], e.get('fila', 0)))

    def clave(e):
        return (_clave_texto(e['categoria']), _clave_texto(e['material']), _clave_texto(e['tipo']),
                _clave_texto(e['marca']), _clave_texto(e['uni']))

    # Orden de aparicion (cronologico): categorias y, dentro de cada una, materiales.
    visto, orden_cat, nombre = {}, {}, {}
    for e in ent:
        k = clave(e)
        if k not in visto:
            visto[k] = len(visto)
            nombre[k] = {'categoria': e['categoria'], 'material': e['material'], 'tipo': e['tipo'],
                         'marca': e['marca'], 'uni': e['uni']}
        orden_cat.setdefault(k[0], len(orden_cat))

    def orden_fila(k):
        return (orden_cat[k[0]], visto[k])

    primera_grafia_s = {}
    primera_grafia_a = {}
    for e in ent:
        s_raw = e.get('suministrador')
        if s_raw is not None:
            sk = _clave_texto(s_raw)
            if sk and sk not in primera_grafia_s:
                primera_grafia_s[sk] = s_raw
        a_raw = e.get('albaran')
        if a_raw is not None:
            ak = _clave_texto(a_raw)
            if ak and ak not in primera_grafia_a:
                primera_grafia_a[ak] = a_raw

    def dict_linea(e):
        s_grafia = primera_grafia_s.get(_clave_texto(e.get('suministrador')))
        a_grafia = primera_grafia_a.get(_clave_texto(e.get('albaran')))
        return {
            'fila': e['fila'],
            'fecha': e['fecha'],
            'suministrador': s_grafia,
            'albaran': a_grafia,
            'completa': entrega_completa(s_grafia, a_grafia),
            'categoria': e.get('categoria'),
            'material': e['material'],
            'tipo': e.get('tipo'),
            'marca': e.get('marca'),
            'uni': e.get('uni'),
            'cantidad': e['cantidad'],
        }

    def clave_orden_linea(e):
        return (
            e['fecha'],
            _clave_texto(e.get('suministrador')),
            _clave_texto(e.get('albaran')),
            orden_fila(clave(e)),
            e.get('fila', 0),
        )

    lineas_todas = [dict_linea(e) for e in sorted(ent, key=clave_orden_linea)]

    meses_claves = sorted({(e['fecha'].year, e['fecha'].month) for e in ent})
    previo = {}          # clave -> cantidad acumulada de meses anteriores
    meses = []
    for (anio, mes) in meses_claves:
        delmes = [e for e in ent if (e['fecha'].year, e['fecha'].month) == (anio, mes)]
        # Una columna por entrega = (fecha, suministrador, albaran), sin distinguir
        # mayusculas ni espacios: "Suelbat" y " SUELBAT " son la misma empresa. Se
        # muestra la primera grafia.
        def col_clave(e):
            return (e['fecha'], _clave_texto(e['suministrador']), _clave_texto(e['albaran']))
        cols_info = {}
        for e in delmes:
            s_grafia = primera_grafia_s.get(_clave_texto(e.get('suministrador')))
            a_grafia = primera_grafia_a.get(_clave_texto(e.get('albaran')))
            cols_info.setdefault(col_clave(e), {'fecha': e['fecha'],
                                                'suministrador': s_grafia,
                                                'albaran': a_grafia,
                                                'completa': entrega_completa(s_grafia, a_grafia)})
        cols = sorted(cols_info)
        filas = {}
        for e in delmes:
            k = clave(e)
            fila = filas.setdefault(k, [None] * len(cols))
            i = cols.index(col_clave(e))
            fila[i] = (fila[i] or 0) + e['cantidad']
        salida = []
        for k in sorted(filas, key=orden_fila):
            cant = [None if x is None else _num(x) for x in filas[k]]
            total = _num(sum(x for x in cant if x is not None))
            ant = _num(previo.get(k, 0))
            salida.append(dict(nombre[k], cantidades=cant, mes_ant=ant, total_mes=total,
                               acumulado=_num(ant + total)))
        for k, cant in filas.items():
            previo[k] = previo.get(k, 0) + sum(x for x in cant if x is not None)
        mes_lineas = [dict_linea(e) for e in sorted(delmes, key=clave_orden_linea)]
        meses.append({'etiqueta': etiqueta_mes(date(anio, mes, 1)), 'anio': anio, 'mes': mes,
                      'columnas': [cols_info[c] for c in cols],
                      'n_incompletas': sum(1 for e in delmes
                                           if not entrega_completa(e['suministrador'], e['albaran'])),
                      'filas': salida,
                      'lineas': mes_lineas})

    etiquetas = [m['etiqueta'] for m in meses]
    acum = []
    for k in sorted(previo, key=orden_fila):
        por_mes = []
        for m in meses:
            f = next((x for x in m['filas'] if _clave_texto(x['categoria']) == k[0]
                      and _clave_texto(x['material']) == k[1] and _clave_texto(x['tipo']) == k[2]
                      and _clave_texto(x['marca']) == k[3] and _clave_texto(x['uni']) == k[4]), None)
            por_mes.append(f['total_mes'] if f else None)
        acum.append(dict(nombre[k], por_mes=por_mes, total=_num(previo[k])))
    return {'meses': meses,
            'acumulado': {'meses': etiquetas, 'filas': acum},
            'ultima_fecha': ent[-1]['fecha'] if ent else None,
            'n_entregas': len(ent),
            'n_incompletas': sum(m['n_incompletas'] for m in meses),
            'lineas': lineas_todas}


def leer_materiales_wb(wb, hoy: date | None = None) -> dict:
    """Mismo diccionario que ``lectores.leer_materiales`` (los consumidores de
    siempre no cambian) mas ``detalle`` y ``formato`` para la vista nueva."""
    hoy = hoy or date.today()
    ent, avisos = leer_registro(wb, hoy)
    g = agrupar(ent)
    res = {'disponible': True, 'meses': g['acumulado']['meses'],
           'ultimo_mes': g['meses'][-1]['etiqueta'] if g['meses'] else None,
           'ultima_fecha': g['ultima_fecha'].strftime('%d/%m/%Y') if g['ultima_fecha'] else None,
           'dias_desde': (hoy - g['ultima_fecha']).days if g['ultima_fecha'] else None,
           'items': [{'categoria': f['categoria'] or '', 'material': f['material'],
                      'tipo': f['tipo'] or '', 'uni': f['uni'] or '', 'total': f['total']}
                     for f in g['acumulado']['filas']],
           'aviso': None, 'formato': 'entregas', 'detalle': g, 'avisos_registro': avisos}
    textos = []
    if avisos:
        textos.append(' · '.join(avisos[:MAX_AVISOS_EN_PANEL])
                      + (f" (+{len(avisos) - MAX_AVISOS_EN_PANEL} mas)" if len(avisos) > MAX_AVISOS_EN_PANEL else ''))
    if res['dias_desde'] is not None and res['dias_desde'] > 30:
        textos.append(f"No hay entregas registradas desde {res['ultima_fecha']} ({res['dias_desde']} dias). "
                      "Los datos de consumo/stock pueden estar desfasados.")
    res['aviso'] = ' '.join(textos) or None
    return res


# ----------------------------------------------------------------------- HTML
#
# Todo el dibujo de las tablas vive en este bloque, a proposito: Bixente quiere
# ver como queda y poder cambiarlo (20/10/2026). Cambiar columnas, orden o
# textos de la vista se hace aqui y en _hoja_mes (Excel), y nada mas.

# Orden por columna al pulsar la cabecera: asc -> desc -> orden original.
# Se inyecta tal cual en la pagina del panel y en el documento del informe.
JS_ORDENAR_TABLA = """function ordenarTabla(th){
  var t=th.closest('table'), tb=t.tBodies[0], col=th.cellIndex;
  var ant=th.getAttribute('data-orden')||'';
  var sig=ant===''?'asc':(ant==='asc'?'desc':'');
  [].forEach.call(t.tHead.querySelectorAll('th'),function(h){
    h.removeAttribute('data-orden'); var f=h.querySelector('.flecha'); if(f){f.textContent='';}
  });
  var filas=[].slice.call(tb.rows);
  if(sig===''){
    filas.sort(function(a,b){return (+a.getAttribute('data-i'))-(+b.getAttribute('data-i'));});
    filas.forEach(function(r){r.style.display='';tb.appendChild(r);});
    return;
  }
  th.setAttribute('data-orden',sig);
  var fl=th.querySelector('.flecha'); if(fl){fl.textContent=sig==='asc'?' \\u25B2':' \\u25BC';}
  function valor(r){var c=r.cells[col]; var v=c?c.getAttribute('data-v'):null; return v===null?(c?c.textContent.trim():''):v;}
  var datos=filas.filter(function(r){return !r.hasAttribute('data-cat');});
  var cats=filas.filter(function(r){return r.hasAttribute('data-cat');});
  cats.forEach(function(r){r.style.display='none';});
  var numerico=datos.every(function(r){var v=valor(r);return v===''||!isNaN(+v);});
  datos.sort(function(a,b){
    var x=valor(a),y=valor(b);
    if(x===''&&y==='')return 0; if(x==='')return 1; if(y==='')return -1;
    var d=numerico?(+x)-(+y):x.localeCompare(y,'es',{numeric:true,sensitivity:'base'});
    return sig==='asc'?d:-d;
  });
  datos.forEach(function(r){tb.appendChild(r);});
  cats.forEach(function(r){tb.appendChild(r);});
}"""

_P_SUBTITULO = 'style="font-size:12.5px;font-weight:700;color:var(--header2);margin:16px 0 6px;"'
_P_META = 'style="font-size:12.5px;color:var(--muted);margin-bottom:10px;"'
_SIN_DATO = '<span style="color:var(--warn)" title="{t}">—</span>'


def _th(texto, derecha=False):
    estilo = 'cursor:pointer;user-select:none;white-space:nowrap;' + ('text-align:right;' if derecha else '')
    return (f'<th style="{estilo}" onclick="ordenarTabla(this)" title="Pulsa para ordenar">'
            f'{_e(texto)}<span class="flecha"></span></th>')


def _td(contenido, valor=None, derecha=False):
    attrs = f' data-v="{_e(valor)}"' if valor is not None else ''
    if derecha:
        attrs += ' style="text-align:right;"'
    return f'<td{attrs}>{contenido}</td>'


def _tabla(cabeceras, filas):
    """Tabla con cabecera pulsable. ``cabeceras``: [(texto, alineada_a_la_derecha)].
    ``filas``: [('cat', texto)] para una fila de categoria (se oculta al ordenar) o
    [('fila', [celdas <td>])]. Cada fila lleva su posicion original (data-i)."""
    cab = ''.join(_th(t, d) for t, d in cabeceras)
    cuerpo = []
    for i, (tipo, contenido) in enumerate(filas):
        if tipo == 'cat':
            cuerpo.append(f'<tr data-i="{i}" data-cat="1"><td colspan="{len(cabeceras)}" '
                          f'style="font-weight:700;color:var(--header2);background:#f4f6f9;">{_e(contenido)}</td></tr>')
        else:
            cuerpo.append(f'<tr data-i="{i}">{"".join(contenido)}</tr>')
    return ('<div style="overflow-x:auto;"><table class="data"><thead><tr>'
            f'{cab}</tr></thead><tbody>{"".join(cuerpo)}</tbody></table></div>')


def _celdas_suministrador_y_albaran(l):
    sumi = _e(l['suministrador']) if l['suministrador'] else _SIN_DATO.format(t='falta el suministrador')
    if l['albaran']:
        alb = _e(l['albaran'])
    elif l['completa']:
        alb = '<span style="color:var(--muted)">en mano</span>'
    else:
        alb = _SIN_DATO.format(t='falta el albarán')
    return (_td(sumi, l['suministrador'] or ''), _td(alb, l['albaran'] or ''))


def _tipo_y_marca(x):
    return ' · '.join(v for v in (x['tipo'], x['marca']) if v)


def _tabla_lineas(lineas):
    """Una fila por linea entregada: Fecha, Suministrador y Albaran como columnas."""
    filas = []
    for l in lineas:
        sumi, alb = _celdas_suministrador_y_albaran(l)
        filas.append(('fila', [_td(f"{l['fecha']:%d/%m/%Y}", f"{l['fecha']:%Y-%m-%d}"), sumi, alb,
                               _td(_e(l['material'])), _td(_e(_tipo_y_marca(l))), _td(_e(l['uni'] or '')),
                               _td(_fmt(l['cantidad']), l['cantidad'], True)]))
    return _tabla([('Fecha', False), ('Suministrador', False), ('Albarán', False), ('Material', False),
                   ('Tipo', False), ('Uni', False), ('Cant.', True)], filas)


def _tabla_totales_mes(m):
    filas, cat = [], object()
    for f in m['filas']:
        if f['categoria'] != cat:
            cat = f['categoria']
            if cat:
                filas.append(('cat', cat))
        filas.append(('fila', [_td(_e(f['material'])), _td(_e(_tipo_y_marca(f))), _td(_e(f['uni'] or '')),
                               _td(_fmt(f['mes_ant']) if f['mes_ant'] else '', f['mes_ant'] or '', True),
                               _td(f'<b>{_fmt(f["total_mes"])}</b>', f['total_mes'], True),
                               _td(_fmt(f['acumulado']), f['acumulado'], True)]))
    return _tabla([('Material', False), ('Tipo', False), ('Uni', False), ('Mes ant.', True),
                   ('Total mes', True), ('Acumulado', True)], filas)


def _vacio():
    return ('<div class="card"><h3>Materiales</h3><p class="empty">Todavia no hay entregas '
            f'registradas en la hoja {HOJA_ENTREGAS}.</p></div>')


def _aviso_incompletas(n):
    return ('<p style="font-size:12.5px;color:var(--warn);margin:0 0 10px;">'
            f'⚠ {n} línea(s) sin suministrador o sin albarán '
            '(entrega en mano: suministrador «Sagarde»). '
            f'Complétalos en la hoja {HOJA_ENTREGAS}.</p>') if n else ''


def render_por_mes(res: dict) -> str:
    """Un bloque por mes (el mas reciente primero): la lista de entregas y debajo
    el total por material."""
    meses = (res.get('detalle') or {}).get('meses') or []
    if not meses:
        return _vacio()
    out = []
    for m in reversed(meses):
        out.append(
            f'<div class="card"><h3>{_e(m["etiqueta"])} — material entregado</h3>'
            f'<p {_P_META}>{len(m["filas"])} material(es) · {len(m["columnas"])} entrega(s) · '
            f'{len(m["lineas"])} línea(s). Cantidades <b>entregadas</b> según el registro de entregas; '
            'no equivalen necesariamente a material instalado.</p>'
            f'{_aviso_incompletas(m.get("n_incompletas"))}'
            f'<div {_P_SUBTITULO}>Entregas del mes</div>{_tabla_lineas(m["lineas"])}'
            f'<div {_P_SUBTITULO}>Total del mes por material</div>{_tabla_totales_mes(m)}</div>')
    return '\n'.join(out)


def render_total(res: dict) -> str:
    """Acumulado de todo lo entregado: una fila por material, una columna por mes."""
    g = res.get('detalle') or {}
    if not g.get('meses'):
        return _vacio()
    ac = g['acumulado']
    filas, cat = [], object()
    for f in ac['filas']:
        if f['categoria'] != cat:
            cat = f['categoria']
            if cat:
                filas.append(('cat', cat))
        filas.append(('fila', [_td(_e(f['material'])), _td(_e(_tipo_y_marca(f))), _td(_e(f['uni'] or ''))]
                      + [_td(_fmt(c), c if c is not None else '', True) for c in f['por_mes']]
                      + [_td(f'<b>{_fmt(f["total"])}</b>', f['total'], True)]))
    cab = [('Material', False), ('Tipo', False), ('Uni', False)] + [(x, True) for x in ac['meses']] + [('Total', True)]
    ultima = g['ultima_fecha'].strftime('%d/%m/%Y') if g.get('ultima_fecha') else '—'
    return ('<div class="card"><h3>Acumulado de material entregado</h3>'
            f'<p {_P_META}>{g["n_entregas"]} línea(s) de entrega · última entrega {ultima}.</p>'
            f'{_tabla(cab, filas)}</div>')


def render_entregas(res: dict) -> str:
    """Todas las entregas, de todos los meses, linea a linea y por orden de fecha."""
    g = res.get('detalle') or {}
    if not g.get('lineas'):
        return _vacio()
    return ('<div class="card"><h3>Todas las entregas, línea a línea</h3>'
            f'<p {_P_META}>{len(g["lineas"])} línea(s), por orden de fecha.</p>'
            f'{_aviso_incompletas(g.get("n_incompletas"))}{_tabla_lineas(g["lineas"])}</div>')


def grupos_por_albaran(lineas):
    """[{'fecha','suministrador','albaran','completa','lineas'}] en orden de fecha."""
    grupos = {}
    for l in lineas:
        k = (l['fecha'], _clave_texto(l['suministrador']), _clave_texto(l['albaran']))
        grupos.setdefault(k, {'fecha': l['fecha'], 'suministrador': l['suministrador'],
                              'albaran': l['albaran'], 'completa': l['completa'], 'lineas': []})['lineas'].append(l)
    return list(grupos.values())


def _titulo_albaran(gr):
    sumi = gr['suministrador'] or 'sin suministrador'
    if gr['albaran']:
        cabeza = f"Albarán {gr['albaran']}"
    elif gr['completa']:
        cabeza = 'Entrega en mano'
    else:
        cabeza = 'Albarán sin número'
    return f"{cabeza} — {sumi} — {gr['fecha']:%d/%m/%Y}"


def render_por_albaran(res: dict) -> str:
    """Un bloque por albaran (fecha + suministrador + numero) con sus materiales."""
    g = res.get('detalle') or {}
    grupos = grupos_por_albaran(g.get('lineas') or [])
    if not grupos:
        return _vacio()
    resumen = []
    for gr in grupos:
        sumi, alb = _celdas_suministrador_y_albaran(dict(gr, lineas=None))
        resumen.append(('fila', [_td(f"{gr['fecha']:%d/%m/%Y}", f"{gr['fecha']:%Y-%m-%d}"), sumi, alb,
                                 _td(str(len(gr['lineas'])), len(gr['lineas']), True)]))
    out = ['<div class="card"><h3>Albaranes</h3>'
           f'<p {_P_META}>{len(grupos)} albarán(es) / entrega(s) en mano.</p>'
           + _tabla([('Fecha', False), ('Suministrador', False), ('Albarán', False), ('Líneas', True)], resumen)
           + '</div>']
    for gr in grupos:
        filas = [('fila', [_td(_e(l['material'])), _td(_e(_tipo_y_marca(l))), _td(_e(l['uni'] or '')),
                           _td(_fmt(l['cantidad']), l['cantidad'], True)]) for l in gr['lineas']]
        falta = '' if gr['completa'] else _aviso_incompletas(len(gr['lineas']))
        out.append(f'<div class="card"><h3>{_e(_titulo_albaran(gr))}</h3>'
                   f'<p {_P_META}>{len(gr["lineas"])} línea(s).</p>{falta}'
                   + _tabla([('Material', False), ('Tipo', False), ('Uni', False), ('Cant.', True)], filas) + '</div>')
    return '\n'.join(out)


def totales_por_suministrador(lineas):
    """[{'suministrador', 'albaranes', 'n_lineas', 'materiales'}] ordenado por
    nombre, con 'sin suministrador' (None) al final. Cada material suma todo lo
    que ese suministrador ha entregado, de todos los meses."""
    por = {}
    for l in lineas:
        k = _clave_texto(l['suministrador'])
        d = por.setdefault(k, {'suministrador': l['suministrador'], 'albaranes': [], 'n_lineas': 0, 'mat': {}})
        d['n_lineas'] += 1
        a = (l['albaran'], l['fecha'])
        if a not in d['albaranes']:
            d['albaranes'].append(a)
        km = (_clave_texto(l['categoria']), _clave_texto(l['material']), _clave_texto(l['tipo']),
              _clave_texto(l['marca']), _clave_texto(l['uni']))
        m = d['mat'].setdefault(km, dict(categoria=l['categoria'], material=l['material'], tipo=l['tipo'],
                                         marca=l['marca'], uni=l['uni'], total=0, orden=len(d['mat'])))
        m['total'] = _num(m['total'] + l['cantidad'])
    salida = []
    for k in sorted(por, key=lambda k: (k == '', k)):
        d = por[k]
        orden_cat = {}
        for m in d['mat'].values():
            orden_cat.setdefault(_clave_texto(m['categoria']), len(orden_cat))
        mats = sorted(d['mat'].values(), key=lambda m: (orden_cat[_clave_texto(m['categoria'])], m['orden']))
        salida.append({'suministrador': d['suministrador'], 'albaranes': d['albaranes'],
                       'n_lineas': d['n_lineas'], 'materiales': mats})
    return salida


def render_por_suministrador(res: dict) -> str:
    """Un bloque por suministrador: sus albaranes y lo que ha entregado en total."""
    g = res.get('detalle') or {}
    grupos = totales_por_suministrador(g.get('lineas') or [])
    if not grupos:
        return _vacio()
    out = []
    for d in grupos:
        nombre = d['suministrador'] or 'Sin suministrador'
        albs = ', '.join((f"{_e(a)} ({f:%d/%m})" if a else f"sin número ({f:%d/%m})") for a, f in d['albaranes'])
        filas, cat = [], object()
        for m in d['materiales']:
            if m['categoria'] != cat:
                cat = m['categoria']
                if cat:
                    filas.append(('cat', cat))
            filas.append(('fila', [_td(_e(m['material'])), _td(_e(_tipo_y_marca(m))), _td(_e(m['uni'] or '')),
                                   _td(f'<b>{_fmt(m["total"])}</b>', m['total'], True)]))
        out.append(f'<div class="card"><h3>{_e(nombre)}</h3>'
                   f'<p {_P_META}>{d["n_lineas"]} línea(s) · Albaranes: {albs}</p>'
                   + _tabla([('Material', False), ('Tipo', False), ('Uni', False), ('Total entregado', True)], filas)
                   + '</div>')
    return '\n'.join(out)


def render_html(res: dict) -> str:
    """Pestana Materiales del panel: por mes y, debajo, el acumulado. Sin
    JavaScript propio ni estilos propios: el mismo bloque se incrusta en el
    informe de obra a la carta (la funcion de ordenar la pone el panel)."""
    if not (res.get('detalle') or {}).get('meses'):
        return _vacio()
    return render_por_mes(res) + '\n' + render_total(res)


# Opciones del "Informe de obra": mismo orden y nombres que los casilleros del panel.
OPCIONES_INFORME = [
    ('total', 'Total acumulado'),
    ('por_mes', 'Por mes'),
    ('por_albaran', 'Por albarán'),
    ('por_suministrador', 'Por suministrador'),
    ('entregas', 'Todas las entregas (línea a línea)'),
]


def secciones_informe(res: dict) -> dict:
    """{opcion: html} para el 'Informe de obra' a la carta. Cada valor se
    calcula con la misma funcion que la pestana: ninguna cifra por un camino aparte."""
    f = {'total': render_total, 'por_mes': render_por_mes, 'por_albaran': render_por_albaran,
         'por_suministrador': render_por_suministrador, 'entregas': render_entregas}
    return {clave: f[clave](res) for clave, _ in OPCIONES_INFORME}


# ----------------------------------------------------------------- Excel

_AZUL = PatternFill('solid', fgColor='0B1F3A') if openpyxl else None
_GRIS = PatternFill('solid', fgColor='E8EDF4') if openpyxl else None
_NOTA = ("Hoja GENERADA automaticamente desde ENTREGAS: no la edites aqui (se sobrescribe). "
         "Anade las filas en ENTREGAS y ejecuta: " + COMANDO_GENERAR)


def _cabecera_estilo(ws, fila, n, desde=1):
    for c in range(desde, desde + n):
        celda = ws.cell(fila, c)
        celda.fill, celda.font = _AZUL, Font(bold=True, color='FFFFFF')
        celda.alignment = Alignment(horizontal='center' if c > desde + 3 else 'left', vertical='center', wrap_text=True)


def _anchos(ws, anchos):
    for i, a in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = a


def _fila_categoria(ws, fila, texto, n, desde=1):
    ws.cell(fila, desde, texto)
    for c in range(desde, desde + n):
        ws.cell(fila, c).fill = _GRIS
        ws.cell(fila, c).font = Font(bold=True)


def _hoja_mes(wb, m):
    ws = wb.create_sheet(m['etiqueta'])
    headers_b1 = ['FECHA', 'SUMINISTRADOR', 'ALBARÁN', 'CATEGORÍA', 'MATERIAL',
                  'TIPO', 'MARCA', 'UNI', 'CANTIDAD']
    ws.append(headers_b1)
    _cabecera_estilo(ws, 1, len(headers_b1))

    lineas = m.get('lineas', [])
    for l in lineas:
        f = l['fecha']
        ws.append([datetime(f.year, f.month, f.day), l['suministrador'], l['albaran'],
                   l['categoria'], l['material'], l['tipo'], l['marca'], l['uni'], l['cantidad']])
        ws.cell(ws.max_row, 1).number_format = 'DD/MM/YYYY'

    ws.auto_filter.ref = f"A1:I{1 + len(lineas)}"

    r_titulo = ws.max_row + 2
    ws.cell(r_titulo, 5, "TOTAL DEL MES POR MATERIAL").font = Font(bold=True)

    r_cab2 = r_titulo + 1
    headers_b2 = ['MATERIAL', 'TIPO', 'MARCA', 'UNI', 'MES ANT', 'TOTAL MES', 'ACUMULADO']
    for col_idx, text in enumerate(headers_b2, start=5):
        ws.cell(r_cab2, col_idx, text)
    _cabecera_estilo(ws, r_cab2, len(headers_b2), desde=5)

    cat = object()
    for f in m['filas']:
        if f['categoria'] != cat:
            cat = f['categoria']
            if cat:
                _fila_categoria(ws, ws.max_row + 1, cat, len(headers_b2), desde=5)
        r_mat = ws.max_row + 1
        vals = [f['material'], f['tipo'], f['marca'], f['uni'], f['mes_ant'] or None, f['total_mes'], f['acumulado']]
        for col_idx, val in enumerate(vals, start=5):
            ws.cell(r_mat, col_idx, val)

    ws.cell(ws.max_row + 2, 1, _NOTA).font = Font(italic=True, color='647184')
    _anchos(ws, [12, 22, 14, 22, 34, 24, 14, 8, 11, 11, 12])
    ws.freeze_panes = 'A2'
    return ws


def _hoja_acumulado(wb, ac):
    ws = wb.create_sheet(HOJA_ACUMULADO)
    n = 4 + len(ac['meses']) + 1
    ws.append(['MATERIAL', 'TIPO', 'MARCA', 'UNI'] + ac['meses'] + ['TOTAL'])
    _cabecera_estilo(ws, 1, n)
    cat = object()
    for f in ac['filas']:
        if f['categoria'] != cat:
            cat = f['categoria']
            if cat:
                _fila_categoria(ws, ws.max_row + 1, cat, n)
        ws.append([f['material'], f['tipo'], f['marca'], f['uni']] + f['por_mes'] + [f['total']])
        ws.cell(ws.max_row, n).font = Font(bold=True)
    ws.cell(ws.max_row + 2, 1, _NOTA).font = Font(italic=True, color='647184')
    _anchos(ws, [34, 22, 14, 8] + [12] * len(ac['meses']) + [12])
    ws.freeze_panes = 'E2'
    return ws


def _hoja_leeme(wb):
    ws = wb.create_sheet(HOJA_LEEME)
    lineas = [
        ("CÓMO USAR ESTE LIBRO", True),
        ("Para añadir material entregado: añade UNA FILA por línea al final de la hoja ENTREGAS "
         "(fecha, suministrador, albarán, categoría, material, tipo, marca, uni, cantidad).", False),
        ("SUMINISTRADOR y ALBARÁN: si lo entregó Sagarde en mano, escribe «Sagarde» y deja el albarán vacío. "
         "Si no, la empresa que figura en el albarán (Electrovizcaya, Suelbat…) y su número de albarán.", False),
        ("Si no sabes el suministrador o el albarán, déjalo vacío: no se inventa. El panel avisará de las líneas "
         "que lo tengan vacío.", False),
        ("Las hojas de cada mes y ACUMULADO se GENERAN a partir de ENTREGAS: no se editan a mano.", False),
        ("Después de añadir filas ejecuta (desde la raíz del repositorio):", False),
        (COMANDO_GENERAR.replace('<obra>', 'olabeaga'), False),
        ("Una fila con fecha ilegible o cantidad no numérica NO se cuenta y se avisa en el panel.", False),
        ("El panel lee ENTREGAS directamente: siempre está al día aunque no hayas regenerado el Excel.", False),
    ]
    for i, (t, negrita) in enumerate(lineas, start=1):
        ws.cell(i, 1, t).font = Font(bold=negrita)
    ws.column_dimensions['A'].width = 130
    return ws


def _validaciones_entregas(ws):
    """Ayudas de entrada de ENTREGAS (idempotente: se rehacen enteras cada vez)."""
    from openpyxl.worksheet.datavalidation import DataValidationList
    ws.data_validations = DataValidationList()
    fechas = DataValidation(type='date', operator='greaterThan', formula1='36526', allow_blank=True,
                            showErrorMessage=True, errorTitle='Fecha', error='Escribe una fecha válida (DD/MM/AAAA).')
    sumi = DataValidation(type='list', formula1='"' + ','.join(SUMINISTRADORES_HABITUALES) + '"',
                          allow_blank=True, showErrorMessage=False, showInputMessage=True,
                          promptTitle='Suministrador',
                          prompt='Entrega en mano: Sagarde. Si no, la empresa del albarán. Puedes escribir otra.')
    alb = DataValidation(type='textLength', operator='greaterThanOrEqual', formula1='0', allow_blank=True,
                         showErrorMessage=False, showInputMessage=True, promptTitle='Albarán',
                         prompt='Número de albarán tal como figura. Vacío si es entrega en mano.')
    cant = DataValidation(type='decimal', operator='notEqual', formula1='0', allow_blank=True,
                          showErrorMessage=True, errorTitle='Cantidad', error='La cantidad debe ser un número distinto de 0.')
    for dv, rango in ((fechas, 'A2:A5000'), (sumi, 'B2:B5000'), (alb, 'C2:C5000'), (cant, 'I2:I5000')):
        ws.add_data_validation(dv)
        dv.add(rango)


def _hoja_entregas(wb, entregas):
    ws = wb.active if wb.sheetnames == ['Sheet'] else wb.create_sheet(HOJA_ENTREGAS, 0)
    ws.title = HOJA_ENTREGAS
    ws.append(CABECERA)
    _cabecera_estilo(ws, 1, len(CABECERA))
    for e in entregas:
        ws.append([datetime(e['fecha'].year, e['fecha'].month, e['fecha'].day), e.get('suministrador'),
                   e.get('albaran'), e.get('categoria'), e['material'], e.get('tipo'), e.get('marca'),
                   e.get('uni'), e['cantidad'], e.get('notas')])
    for fila in ws.iter_rows(min_row=2, max_col=1):
        fila[0].number_format = 'DD/MM/YYYY'
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f"A1:{get_column_letter(len(CABECERA))}{max(ws.max_row, 2)}"
    _anchos(ws, [12, 22, 14, 22, 30, 24, 14, 8, 11, 30])
    _validaciones_entregas(ws)
    return ws


def construir_hojas(wb, entregas) -> dict:
    """Borra las hojas generadas (meses y ACUMULADO) y las rehace desde las
    entregas. No toca ENTREGAS ni ninguna hoja ajena. Devuelve la agrupacion."""
    for nombre in list(wb.sheetnames):
        if _RE_HOJA_MES.match(nombre) or nombre == HOJA_ACUMULADO:
            del wb[nombre]
    g = agrupar(entregas)
    for m in g['meses']:
        _hoja_mes(wb, m)
    _hoja_acumulado(wb, g['acumulado'])
    if HOJA_LEEME in wb.sheetnames:       # tambien es texto generado: se mantiene al dia
        del wb[HOJA_LEEME]
    _hoja_leeme(wb)
    if HOJA_ENTREGAS in wb.sheetnames:
        _validaciones_entregas(wb[HOJA_ENTREGAS])
    orden = {nombre: (1, i) for i, nombre in enumerate(m['etiqueta'] for m in g['meses'])}

    def pos(ws):
        if ws.title == HOJA_ENTREGAS:
            return (0, 0)
        if ws.title in orden:
            return (1,) + orden[ws.title][1:]
        if ws.title == HOJA_ACUMULADO:
            return (2, 0)
        if ws.title == HOJA_LEEME:
            return (3, 0)
        return (4, 0)
    wb._sheets.sort(key=pos)
    wb.active = 0
    return g


def _guardar(wb, ruta):
    tmp = ruta + '.generando.tmp'
    wb.save(tmp)
    os.replace(tmp, ruta)


def crear_libro(entregas, ruta):
    """Crea el libro completo (ENTREGAS + meses + ACUMULADO + LEEME)."""
    wb = openpyxl.Workbook()
    _hoja_entregas(wb, entregas)
    construir_hojas(wb, entregas)
    _guardar(wb, ruta)


def _abrir(ruta, data_only=False):
    with open(ruta, 'rb') as f:        # desde memoria: no deja el fichero retenido (Windows)
        return openpyxl.load_workbook(io.BytesIO(f.read()), data_only=data_only)


def regenerar_hojas(ruta, hoy: date | None = None):
    """Rehace las hojas generadas de un libro existente. Devuelve los avisos."""
    wb = _abrir(ruta)
    entregas, avisos = leer_registro(wb, hoy)
    construir_hojas(wb, entregas)
    _guardar(wb, ruta)
    return avisos


def refrescar_hojas(ruta, hoy: date | None = None) -> bool:
    """Si ``ruta`` es un registro de entregas y sus hojas generadas se han quedado
    atras respecto a ENTREGAS, las rehace (con copia previa en ``_SISTEMA``).
    Devuelve True si las rehizo. Existe para que el resultado no dependa de que
    quien edita ENTREGAS (agy, Claude, una persona) se acuerde del ultimo paso:
    lo llama ``generar_todos`` cada vez que regenera la obra."""
    if not os.path.isfile(ruta):
        return False
    wb = _abrir(ruta, data_only=True)
    if HOJA_ENTREGAS not in wb.sheetnames or not hojas_desactualizadas(wb):
        return False
    carpeta = os.path.join(os.path.dirname(os.path.abspath(ruta)), '_SISTEMA')
    os.makedirs(carpeta, exist_ok=True)
    marca = datetime.now().strftime('%Y%m%d_%H%M%S')
    shutil.copy2(ruta, os.path.join(carpeta, f"{os.path.basename(ruta)}.antes_refrescar_{marca}.bak"))
    regenerar_hojas(ruta, hoy)
    return True


def _grid_recortado(ws):
    filas = [list(r) for r in ws.iter_rows(values_only=True)]
    while filas and all(x is None for x in filas[-1]):
        filas.pop()
    ancho = max((max((i for i, x in enumerate(r) if x is not None), default=-1) + 1 for r in filas), default=0)
    return [r[:ancho] + [None] * (ancho - len(r[:ancho])) for r in filas]


def hojas_desactualizadas(wb) -> list[str]:
    """Compara las hojas generadas del libro con las que saldrian hoy de ENTREGAS."""
    entregas, _ = leer_registro(wb)
    esperado = openpyxl.Workbook()
    esperado.remove(esperado.active)
    construir_hojas(esperado, entregas)
    problemas = []
    nombres_esperados = [n for n in esperado.sheetnames if n != HOJA_LEEME]
    for n in nombres_esperados:
        if n not in wb.sheetnames:
            problemas.append(f"falta la hoja {n}")
        elif _grid_recortado(wb[n]) != _grid_recortado(esperado[n]):
            problemas.append(f"la hoja {n} no coincide con ENTREGAS")
    for n in wb.sheetnames:
        if (_RE_HOJA_MES.match(n) or n == HOJA_ACUMULADO) and n not in nombres_esperados:
            problemas.append(f"sobra la hoja {n}")
    return problemas


# --------------------------------------------------------------- migracion

def entregas_desde_rejilla(wb) -> list[dict]:
    """Convierte la rejilla antigua (catalogo x fechas) en filas de entrega.
    Solo cuentan las columnas con fecha: MES ANT y TOTAL son arrastres/derivados
    y no se migran. Suministrador y albaran quedan sin dato (no existian)."""
    entregas = []
    for ws in wb.worksheets:
        filas = list(ws.iter_rows(values_only=True))
        if not filas:
            continue
        fechas = {i: _fecha(h) for i, h in enumerate(filas[0]) if isinstance(h, (datetime, date))}
        if not fechas:
            continue
        cat = None
        for fila in filas[1:]:
            nombre = fila[0] if fila else None
            if not isinstance(nombre, str) or not nombre.strip():
                continue
            cantidades = [(i, fila[i]) for i in fechas
                          if i < len(fila) and isinstance(fila[i], (int, float))
                          and not isinstance(fila[i], bool) and fila[i]]
            tipo = fila[1] if len(fila) > 1 else None
            if not cantidades and tipo in (None, ''):
                cat = nombre.strip()
                continue
            for i, c in cantidades:
                entregas.append({'fila': len(entregas) + 2, 'fecha': fechas[i], 'suministrador': None,
                                 'albaran': None, 'categoria': cat, 'material': ' '.join(nombre.split()),
                                 'tipo': _txt(tipo), 'marca': _txt(fila[2] if len(fila) > 2 else None),
                                 'uni': _txt(fila[3] if len(fila) > 3 else None), 'cantidad': _num(c),
                                 'notas': None})
    entregas.sort(key=lambda e: e['fecha'])      # estable: conserva el orden del catalogo dentro de cada dia
    for n, e in enumerate(entregas, start=2):
        e['fila'] = n
    return entregas
