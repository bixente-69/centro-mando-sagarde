# -*- coding: utf-8 -*-
"""Pruebas de la hoja de entrega de materiales: lector del panel y herramienta
`materiales_xlsx.py` (comprobar / reparar).

Nacen del 20/10/2026: agy anadio material a Olabeaga editando a mano el XML del
Excel. Dejo dos defectos que no daban error visible:

  1. ``<?xml ... encoding="utf-16"?>`` sobre bytes UTF-8 -> openpyxl no abre
     el fichero.
  2. 505 celdas de formula con ``<v>`` solo de espacios -> openpyxl revienta
     con ``ValueError`` al leer el valor.

``leer_materiales`` lo propagaba como excepcion, ``generar_todos.main()`` la
atrapaba dentro del bloque de TODO el panel y descartaba la obra entera: el
indice y el portal movil dejaron de mostrar Olabeaga sin ningun aviso.

Contrato que fijan estas pruebas:
  - Un Excel de materiales malo NUNCA tumba el panel: el lector devuelve un
    ``aviso`` visible y ``items`` vacio (no se inventan totales).
  - Un Excel guardado sin recalcular (formulas sin valor en cache) tampoco
    se lee en silencio como "sin materiales": avisa.
  - La herramienta detecta cada defecto y repara los de XML sin perder datos.
"""
import io
import os
import re
import sys
import tempfile
import unittest
import zipfile
from contextlib import redirect_stdout
from datetime import datetime

import openpyxl

_SISTEMA_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAIZ = os.path.normpath(os.path.join(_SISTEMA_DIR, "..", ".."))
SCRIPTS = os.path.join(RAIZ, "_SISTEMA", "MOTOR", "scripts")
for _p in (_SISTEMA_DIR, os.path.join(_SISTEMA_DIR, "adaptadores"), SCRIPTS):
    if _p not in sys.path:
        sys.path.insert(0, _p)

import lectores  # noqa: E402
import materiales_xlsx as mx  # noqa: E402

HOJA_XML = "xl/worksheets/sheet1.xml"


def _crear_sano(ruta, con_formulas_sin_cache=False):
    """Hoja minima con la misma forma que las de las obras: cabecera con
    MATERIAL/TIPO/MARCA/UNI, una fecha, TOTAL y una fila de categoria."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Octubre26"
    ws.append(["MATERIAL", "TIPO", "MARCA", "UNI",
               datetime.now().replace(hour=0, minute=0, second=0, microsecond=0),
               "TOTAL"])
    ws.append(["SUJECCION Y MARCADO", None, None, None, None, None])
    ws.append(["BRIDAS", "3,6x280", None, 100, 12, "=SUM(E3:E3)" if con_formulas_sin_cache else 12])
    wb.save(ruta)


def _reescribir_hoja(ruta, transformar):
    """Reescribe ``sheet1.xml`` aplicando ``transformar(bytes) -> bytes``."""
    tmp = ruta + ".tmp"
    with zipfile.ZipFile(ruta) as zin, zipfile.ZipFile(tmp, "w") as zout:
        for info in zin.infolist():
            datos = zin.read(info.filename)
            if info.filename == HOJA_XML:
                datos = transformar(datos)
            zout.writestr(info, datos, compress_type=info.compress_type)
    os.replace(tmp, ruta)


def _declaracion_utf16_mal(datos):
    """Defecto 1 real de agy: declara utf-16 sobre bytes UTF-8."""
    datos = re.sub(rb"^<\?xml[^>]*\?>\s*", b"", datos)
    return b'<?xml version="1.0" encoding="utf-16"?>\r\n' + datos


def _v_en_blanco(datos):
    """Defecto 2 real de agy: una formula con <v> de solo espacios."""
    celda = b'<c r="Z2"><f>A2</f><v>\r\n        </v></c>'
    return datos.replace(b"</row>", celda + b"</row>", 1)


def _total(res):
    return sorted(i["total"] for i in res["items"])


def _celdas(ruta):
    """Valores de la primera hoja, cargada desde memoria para no dejar el
    fichero retenido (en Windows impide borrar el directorio temporal)."""
    with open(ruta, "rb") as f:
        ws = openpyxl.load_workbook(io.BytesIO(f.read())).active
    return [[c.value for c in fila] for fila in ws.iter_rows()]


class TestLectorTolerante(unittest.TestCase):
    """leer_materiales no lanza nunca por un Excel danado."""

    def _ruta(self, tmp, nombre="m.xlsx"):
        return os.path.join(tmp, nombre)

    def test_sano_devuelve_los_items(self):
        with tempfile.TemporaryDirectory() as tmp:
            r = self._ruta(tmp)
            _crear_sano(r)
            res = lectores.leer_materiales(r)
        self.assertTrue(res["disponible"])
        self.assertEqual(_total(res), [12])
        self.assertIsNone(res["aviso"])

    def test_declaracion_de_codificacion_mala_no_tumba_el_lector(self):
        with tempfile.TemporaryDirectory() as tmp:
            r = self._ruta(tmp)
            _crear_sano(r)
            _reescribir_hoja(r, _declaracion_utf16_mal)
            res = lectores.leer_materiales(r)   # no debe lanzar
        self.assertTrue(res["disponible"])
        self.assertEqual(res["items"], [])
        self.assertIn("no se pudo abrir", res["aviso"])
        self.assertIn("materiales_xlsx.py", res["aviso"])

    def test_celdas_con_valor_en_blanco_no_tumban_el_lector(self):
        with tempfile.TemporaryDirectory() as tmp:
            r = self._ruta(tmp)
            _crear_sano(r)
            _reescribir_hoja(r, _v_en_blanco)
            res = lectores.leer_materiales(r)   # no debe lanzar
        self.assertTrue(res["disponible"])
        self.assertEqual(res["items"], [])
        self.assertIn("no se pudo abrir", res["aviso"])

    def test_formulas_sin_valor_calculado_avisan_en_vez_de_callar(self):
        """Un Excel guardado sin recalcular tiene el TOTAL vacio en cache:
        sin aviso el panel diria 'sin materiales' y nadie lo notaria."""
        with tempfile.TemporaryDirectory() as tmp:
            r = self._ruta(tmp)
            _crear_sano(r, con_formulas_sin_cache=True)
            res = lectores.leer_materiales(r)
        self.assertTrue(res["disponible"])
        self.assertEqual(res["items"], [])
        self.assertIn("sin calcular", res["aviso"])
        self.assertIn("materiales_xlsx.py", res["aviso"])


class TestDiagnostico(unittest.TestCase):
    def test_fichero_sano_no_tiene_defectos(self):
        with tempfile.TemporaryDirectory() as tmp:
            r = os.path.join(tmp, "m.xlsx")
            _crear_sano(r)
            self.assertEqual(mx.diagnosticar(r), [])

    def test_detecta_codificacion_declarada_que_no_es_la_real(self):
        with tempfile.TemporaryDirectory() as tmp:
            r = os.path.join(tmp, "m.xlsx")
            _crear_sano(r)
            _reescribir_hoja(r, _declaracion_utf16_mal)
            defectos = mx.diagnosticar(r)
        self.assertTrue(any("codific" in d for d in defectos), defectos)

    def test_detecta_valores_en_blanco(self):
        with tempfile.TemporaryDirectory() as tmp:
            r = os.path.join(tmp, "m.xlsx")
            _crear_sano(r)
            _reescribir_hoja(r, _v_en_blanco)
            defectos = mx.diagnosticar(r)
        self.assertTrue(any("en blanco" in d for d in defectos), defectos)

    def test_detecta_formulas_sin_valor_calculado(self):
        with tempfile.TemporaryDirectory() as tmp:
            r = os.path.join(tmp, "m.xlsx")
            _crear_sano(r, con_formulas_sin_cache=True)
            defectos = mx.diagnosticar(r)
        self.assertTrue(any("sin valor calculado" in d for d in defectos), defectos)

    def test_fichero_inexistente_y_no_zip_se_dicen(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertTrue(mx.diagnosticar(os.path.join(tmp, "no_existe.xlsx")))
            falso = os.path.join(tmp, "falso.xlsx")
            with open(falso, "wb") as f:
                f.write(b"esto no es un zip")
            self.assertTrue(any("zip" in d.lower() or "xlsx" in d.lower()
                                for d in mx.diagnosticar(falso)))


class TestReparar(unittest.TestCase):
    def test_reparar_xml_arregla_los_dos_defectos_sin_perder_datos(self):
        with tempfile.TemporaryDirectory() as tmp:
            r = os.path.join(tmp, "m.xlsx")
            _crear_sano(r)
            antes = _celdas(r)
            _reescribir_hoja(r, _declaracion_utf16_mal)
            _reescribir_hoja(r, _v_en_blanco)
            acciones = mx.reparar_xml(r)
            despues = _celdas(r)
            # openpyxl en modo formulas ignora los <v>: lo que cuenta es que el
            # diagnostico ya no vea defectos de XML y que la lectura que hace
            # el panel (data_only) funcione y devuelva los datos.
            restantes = [d for d in mx.diagnosticar(r)
                         if "en blanco" in d or "codific" in d]
            lectura = lectores.leer_materiales(r)
        self.assertTrue(acciones)
        self.assertEqual(restantes, [])
        self.assertEqual(_total(lectura), [12])
        self.assertIsNone(lectura["aviso"])
        # la celda de formula anadida por el defecto es lo unico que sobra
        self.assertEqual([f[:6] for f in despues], [f[:6] for f in antes])

    def test_reparar_hace_copia_en_sistema_y_recalcula_si_hace_falta(self):
        with tempfile.TemporaryDirectory() as tmp:
            r = os.path.join(tmp, "REVISIONES", "m.xlsx")
            os.makedirs(os.path.dirname(r))
            _crear_sano(r, con_formulas_sin_cache=True)
            llamadas = []

            def recalculo_falso(ruta):
                llamadas.append(ruta)
                _crear_sano(ruta)   # deja el fichero con valores en cache

            res = mx.reparar(r, recalcular=recalculo_falso)
            copias = os.listdir(os.path.join(tmp, "REVISIONES", "_SISTEMA"))
        self.assertTrue(res["ok"], res)
        self.assertEqual(len(llamadas), 1)
        self.assertEqual(len(copias), 1)
        self.assertIn("antes_reparar", copias[0])

    def test_reparar_un_fichero_sano_no_toca_nada(self):
        with tempfile.TemporaryDirectory() as tmp:
            r = os.path.join(tmp, "REVISIONES", "m.xlsx")
            os.makedirs(os.path.dirname(r))
            _crear_sano(r)
            res = mx.reparar(r, recalcular=lambda ruta: self.fail("no debia recalcular"))
            self.assertTrue(res["ok"])
            self.assertIsNone(res["copia"])
            self.assertFalse(os.path.exists(os.path.join(tmp, "REVISIONES", "_SISTEMA")))

    def test_sin_excel_no_se_declara_arreglado_lo_que_no_lo_esta(self):
        """Sin Excel no hay recalculo: ok=False y se dice que hay que abrir
        el fichero en Excel y guardar. Nunca un 'todo bien' falso."""
        def sin_excel(ruta):
            raise mx.RecalculoNoDisponible("Excel no esta instalado")

        with tempfile.TemporaryDirectory() as tmp:
            r = os.path.join(tmp, "REVISIONES", "m.xlsx")
            os.makedirs(os.path.dirname(r))
            _crear_sano(r, con_formulas_sin_cache=True)
            res = mx.reparar(r, recalcular=sin_excel)
        self.assertFalse(res["ok"])
        self.assertTrue(any("Excel" in p for p in res["pendientes"]), res)

    def test_main_comprobar_devuelve_codigo_de_salida(self):
        with tempfile.TemporaryDirectory() as tmp:
            sano = os.path.join(tmp, "sano.xlsx")
            roto = os.path.join(tmp, "roto.xlsx")
            _crear_sano(sano)
            _crear_sano(roto)
            _reescribir_hoja(roto, _declaracion_utf16_mal)
            with redirect_stdout(io.StringIO()):
                self.assertEqual(mx.main(["comprobar", sano]), 0)
                self.assertEqual(mx.main(["comprobar", roto]), 1)

    def test_ruta_por_id_de_obra_apunta_a_su_hoja(self):
        ruta = mx.ruta_materiales("olabeaga")
        self.assertTrue(ruta.endswith(
            os.path.join("2026 OLABEAGA BILBAO", "REVISIONES",
                         "hoja de entrega de materiales OLABEAGA.xlsx")), ruta)


if __name__ == "__main__":
    unittest.main()
