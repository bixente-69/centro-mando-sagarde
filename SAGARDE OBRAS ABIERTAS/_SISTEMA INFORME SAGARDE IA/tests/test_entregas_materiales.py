# -*- coding: utf-8 -*-
"""Registro de entregas de material (piloto Olabeaga, 20/10/2026).

Antes: un Excel con el catalogo entero x una columna por fecha, sin sitio para
suministrador ni albaran, que el panel leia solo en su ultima hoja. Despues: una
hoja ENTREGAS (una fila por linea entregada) es la UNICA fuente; las hojas de
mes, el acumulado y la pestana del panel se generan de ahi, asi que no pueden
divergir.

Contrato que fijan estas pruebas:
  - Una fila invalida NO se cuenta ni se convierte en cero: queda fuera y se
    avisa con su numero de fila.
  - Suministrador o albaran desconocidos son 'sin dato' (None), nunca "0".
  - Las hojas generadas llevan solo valores (sin formulas): no hay totales
    "sin calcular" posibles.
  - Regenerar es idempotente y no toca ENTREGAS ni hojas ajenas.
"""
import io
import os
import sys
import tempfile
import unittest
from datetime import date, datetime

import openpyxl

_SISTEMA_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAIZ = os.path.normpath(os.path.join(_SISTEMA_DIR, "..", ".."))
SCRIPTS = os.path.join(RAIZ, "_SISTEMA", "MOTOR", "scripts")
for _p in (_SISTEMA_DIR, os.path.join(_SISTEMA_DIR, "adaptadores"), SCRIPTS):
    if _p not in sys.path:
        sys.path.insert(0, _p)

import entregas_materiales as em  # noqa: E402
import lectores  # noqa: E402
import materiales_xlsx as mx  # noqa: E402
import panel_obra  # noqa: E402

D = date
HOY = D(2027, 1, 31)   # reloj fijo: la muestra incluye un 03/11/2026 que hoy aun es futuro


def _fila(fecha, sumi, alb, cat, mat, tipo, marca, uni, cant, notas=None):
    return [fecha, sumi, alb, cat, mat, tipo, marca, uni, cant, notas]


MUESTRA = [
    _fila(datetime(2026, 10, 1), "Sumi A", "111", "CABLE 1,5", "AZUL", 1.5, None, 200, 400),
    _fila(datetime(2026, 10, 1), "Sumi A", "111", "CABLE 1,5", "GRIS", 1.5, None, 200, 400),
    _fila(datetime(2026, 10, 5), "Sumi B", "222", "TUBO", "PVC", "GRIS", None, 20, 171),
    _fila(datetime(2026, 10, 5), "Sumi B", "223", "TUBO", "PVC", "GRIS", None, 20, 29),
    _fila(datetime(2026, 11, 3), "Sumi A", "333", "CABLE 1,5", "AZUL", 1.5, None, 200, 600),
]


def _wb(filas=None, cabecera=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ENTREGAS"
    ws.append(cabecera or em.CABECERA)
    for f in (MUESTRA if filas is None else filas):
        ws.append(f)
    return wb


def _guardar(wb, ruta):
    wb.save(ruta)
    return ruta


def _cargar(ruta, valores=True):
    with open(ruta, "rb") as f:
        return openpyxl.load_workbook(io.BytesIO(f.read()), data_only=valores)


def _grid(ws):
    return [[c.value for c in fila] for fila in ws.iter_rows()]


class TestLecturaDelRegistro(unittest.TestCase):

    def test_lee_una_fila_valida_con_todos_sus_campos(self):
        ent, av = em.leer_registro(_wb(MUESTRA[:1]))
        self.assertEqual(av, [])
        self.assertEqual(len(ent), 1)
        e = ent[0]
        self.assertEqual((e["fecha"], e["suministrador"], e["albaran"]),
                         (D(2026, 10, 1), "Sumi A", "111"))
        self.assertEqual((e["categoria"], e["material"], e["uni"], e["cantidad"]),
                         ("CABLE 1,5", "AZUL", "200", 400))

    def test_fecha_en_texto_dd_mm_aaaa_y_albaran_numerico(self):
        fila = _fila("20/10/2026", "X", 4567, "TUBO", "PVC", "GRIS", None, 20, 5)
        ent, av = em.leer_registro(_wb([fila]))
        self.assertEqual(av, [])
        self.assertEqual(ent[0]["fecha"], D(2026, 10, 20))
        self.assertEqual(ent[0]["albaran"], "4567")   # sin ".0" de coma flotante

    def test_cabecera_con_tildes_y_otro_orden(self):
        cab = ["Cantidad", "Material", "Fecha", "Albarán", "Suministrador"]
        wb = _wb([[10, "BRIDAS", datetime(2026, 10, 2), "A-9", "Sumi"]], cabecera=cab)
        ent, av = em.leer_registro(wb)
        self.assertEqual(av, [])
        self.assertEqual((ent[0]["material"], ent[0]["cantidad"], ent[0]["albaran"]),
                         ("BRIDAS", 10, "A-9"))

    def test_falta_una_columna_obligatoria_se_dice_y_no_se_lee_nada(self):
        wb = _wb([], cabecera=["FECHA", "MATERIAL"])
        ent, av = em.leer_registro(wb)
        self.assertEqual(ent, [])
        self.assertTrue(any("CANTIDAD" in a for a in av), av)

    def test_cantidad_no_numerica_queda_fuera_y_no_es_un_cero(self):
        malas = [_fila(datetime(2026, 10, 1), "S", "1", "T", "PVC", None, None, 20, "unas pocas")]
        ent, av = em.leer_registro(_wb(MUESTRA[:1] + malas))
        self.assertEqual(len(ent), 1)
        self.assertTrue(any("fila 3" in a and "cantidad" in a.lower() for a in av), av)

    def test_fecha_invalida_queda_fuera_con_aviso(self):
        malas = [_fila("ayer", "S", "1", "T", "PVC", None, None, 20, 5)]
        ent, av = em.leer_registro(_wb(malas))
        self.assertEqual(ent, [])
        self.assertTrue(any("fila 2" in a and "fecha" in a.lower() for a in av), av)

    def test_cantidad_en_texto_ambigua_se_rechaza_y_la_clara_se_acepta(self):
        """'1.200' puede ser 1,2 o 1200: no se adivina. '1200,5' y '12' valen."""
        for texto in ("1.200", "2,500"):
            ent, av = em.leer_registro(_wb([_fila(datetime(2026, 10, 1), "S", "1", "T", "PVC", None, None, 20, texto)]))
            self.assertEqual(ent, [], texto)
            self.assertTrue(any("ambigua" in a for a in av), (texto, av))
        ent, av = em.leer_registro(_wb([_fila(datetime(2026, 10, 1), "S", "1", "T", "PVC", None, None, 20, "1200,5"),
                                        _fila(datetime(2026, 10, 2), "S", "2", "T", "PVC", None, None, 20, "12")]))
        self.assertEqual([e["cantidad"] for e in ent], [1200.5, 12])

    def test_cantidad_cero_se_ignora_con_aviso(self):
        ent, av = em.leer_registro(_wb([_fila(datetime(2026, 10, 1), "S", "1", "T", "PVC", None, None, 20, 0)]))
        self.assertEqual(ent, [])
        self.assertTrue(any("fila 2" in a for a in av), av)

    def test_misma_linea_repetida_se_cuenta_pero_se_avisa(self):
        """Dos filas identicas (fecha, albaran, material...) suelen ser una
        entrada doble de agy: se avisa, y no se descarta en silencio porque
        podrian ser dos entregas reales sin albaran."""
        dup = MUESTRA[:1] + MUESTRA[:1]
        ent, av = em.leer_registro(_wb(dup))
        self.assertEqual(len(ent), 2)
        self.assertTrue(any("repetida" in a.lower() and "2" in a and "3" in a for a in av), av)

    def test_fecha_futura_se_avisa(self):
        futura = _fila(datetime(2999, 1, 1), "S", "1", "T", "PVC", None, None, 20, 5)
        ent, av = em.leer_registro(_wb([futura]))
        self.assertEqual(len(ent), 1)
        self.assertTrue(any("futura" in a.lower() for a in av), av)

    def test_filas_vacias_se_saltan_sin_avisar(self):
        ent, av = em.leer_registro(_wb(MUESTRA[:1] + [[None] * 10] + MUESTRA[1:2]))
        self.assertEqual((len(ent), av), (2, []))


class TestAgrupacion(unittest.TestCase):

    def setUp(self):
        ent, _ = em.leer_registro(_wb())
        self.g = em.agrupar(ent)
        self.oct, self.nov = self.g["meses"]

    def test_meses_en_orden_cronologico_con_etiqueta_como_las_hojas_de_siempre(self):
        self.assertEqual([m["etiqueta"] for m in self.g["meses"]], ["Octubre26", "Noviembre26"])

    def test_misma_fecha_distinto_albaran_son_dos_columnas(self):
        cols = [(c["fecha"], c["suministrador"], c["albaran"]) for c in self.oct["columnas"]]
        self.assertEqual(cols, [(D(2026, 10, 1), "Sumi A", "111"),
                                (D(2026, 10, 5), "Sumi B", "222"),
                                (D(2026, 10, 5), "Sumi B", "223")])

    def test_totales_y_arrastre_entre_meses(self):
        azul_oct = next(f for f in self.oct["filas"] if f["material"] == "AZUL")
        azul_nov = next(f for f in self.nov["filas"] if f["material"] == "AZUL")
        self.assertEqual((azul_oct["mes_ant"], azul_oct["total_mes"], azul_oct["acumulado"]), (0, 400, 400))
        self.assertEqual((azul_nov["mes_ant"], azul_nov["total_mes"], azul_nov["acumulado"]), (400, 600, 1000))

    def test_pvc_suma_las_dos_entregas_del_mismo_dia(self):
        pvc = next(f for f in self.oct["filas"] if f["material"] == "PVC")
        self.assertEqual(pvc["cantidades"], [None, 171, 29])
        self.assertEqual(pvc["total_mes"], 200)

    def test_un_mes_solo_lista_lo_entregado_ese_mes_sin_repetir_el_catalogo(self):
        self.assertEqual(sorted(f["material"] for f in self.nov["filas"]), ["AZUL"])

    def test_acumulado_recoge_todo_lo_entregado_por_mes_y_total(self):
        ac = {f["material"]: f for f in self.g["acumulado"]["filas"]}
        self.assertEqual(self.g["acumulado"]["meses"], ["Octubre26", "Noviembre26"])
        self.assertEqual(ac["AZUL"]["por_mes"], [400, 600])
        self.assertEqual(ac["AZUL"]["total"], 1000)
        self.assertEqual(ac["GRIS"]["total"], 400)

    def test_mayusculas_y_espacios_distintos_son_la_misma_fila(self):
        filas = [_fila(datetime(2026, 10, 1), "S", "1", "CABLE", "NEGRO ", 1.5, None, 200, 100),
                 _fila(datetime(2026, 10, 2), "S", "2", "CABLE", "negro", 1.5, None, 200, 50)]
        g = em.agrupar(em.leer_registro(_wb(filas))[0])
        self.assertEqual(len(g["meses"][0]["filas"]), 1)
        self.assertEqual(g["meses"][0]["filas"][0]["total_mes"], 150)

    def test_mismo_material_distinto_tipo_son_filas_distintas(self):
        filas = [_fila(datetime(2026, 10, 1), "S", "1", "CABLE", "AZUL", 1.5, None, 200, 100),
                 _fila(datetime(2026, 10, 1), "S", "1", "CABLE", "AZUL", 2.5, None, 200, 50)]
        g = em.agrupar(em.leer_registro(_wb(filas))[0])
        self.assertEqual(len(g["meses"][0]["filas"]), 2)

    def test_suministrador_y_albaran_desconocidos_son_sin_dato_no_cero(self):
        filas = [_fila(datetime(2026, 10, 1), None, None, "T", "PVC", None, None, 20, 5)]
        g = em.agrupar(em.leer_registro(_wb(filas))[0])
        col = g["meses"][0]["columnas"][0]
        self.assertIsNone(col["suministrador"])
        self.assertIsNone(col["albaran"])


class TestCompatibilidadConElLectorDeSiempre(unittest.TestCase):

    def test_resumen_con_las_mismas_claves_que_leer_materiales(self):
        res = em.leer_materiales_wb(_wb(), hoy=HOY)
        for k in ("disponible", "meses", "ultimo_mes", "ultima_fecha", "dias_desde", "items", "aviso"):
            self.assertIn(k, res)
        self.assertEqual(res["meses"], ["Octubre26", "Noviembre26"])
        self.assertEqual(res["ultimo_mes"], "Noviembre26")
        self.assertEqual(res["ultima_fecha"], "03/11/2026")
        azul = next(i for i in res["items"] if i["material"] == "AZUL")
        self.assertEqual(azul["total"], 1000)           # acumulado a fecha, como el TOTAL de siempre
        self.assertEqual(set(azul), {"categoria", "material", "tipo", "uni", "total"})

    def test_leer_materiales_despacha_por_la_hoja_entregas(self):
        with tempfile.TemporaryDirectory() as tmp:
            ruta = _guardar(_wb(MUESTRA[:4]), os.path.join(tmp, "m.xlsx"))
            res = lectores.leer_materiales(ruta)
        self.assertEqual(res["ultimo_mes"], "Octubre26")
        self.assertIn("detalle", res)
        self.assertEqual(len(res["items"]), 3)    # AZUL, GRIS y PVC

    def test_los_avisos_de_filas_malas_llegan_al_panel(self):
        malas = MUESTRA[:1] + [_fila("ayer", "S", "1", "T", "PVC", None, None, 20, 5)]
        with tempfile.TemporaryDirectory() as tmp:
            ruta = _guardar(_wb(malas), os.path.join(tmp, "m.xlsx"))
            res = lectores.leer_materiales(ruta)
        self.assertIn("fila 3", res["aviso"])


class TestPanel(unittest.TestCase):

    def _html(self, filas=None):
        res = em.leer_materiales_wb(_wb(filas), hoy=HOY)
        return em.render_html(res)

    def test_muestra_mes_fechas_suministrador_albaran_y_totales(self):
        h = self._html()
        for esperado in ("Octubre26", "Noviembre26", "Sumi A", "Sumi B", "111", "222", "223",
                         "AZUL", "PVC", "01/10", "Acumulado"):
            self.assertIn(esperado, h)
        self.assertIn("1.000", h)           # acumulado de AZUL con separador de miles del panel

    def test_todas_las_filas_de_la_tabla_de_un_mes_ocupan_las_mismas_columnas(self):
        """Fallo real de la primera version: las filas Suministrador/Albaran
        llevaban colspan=5 y quedaban desplazadas una columna, de modo que el
        albaran de una entrega aparecia bajo la fecha de la siguiente."""
        import re
        h = self._html(MUESTRA[:4])
        tabla = h[h.index("<table"):h.index("</table>")]
        anchos = []
        for tr in re.findall(r"<tr>(.*?)</tr>", tabla, re.S):
            ancho = 0
            for celda in re.findall(r"<t[hd]\b([^>]*)>", tr):
                m = re.search(r'colspan="(\d+)"', celda)
                ancho += int(m.group(1)) if m else 1
            anchos.append(ancho)
        self.assertEqual(len(set(anchos)), 1, anchos)

    def test_el_albaran_queda_bajo_la_fecha_de_su_entrega(self):
        import re
        h = self._html(MUESTRA[:4])
        tabla = h[h.index("<table"):h.index("</table>")]
        filas = re.findall(r"<tr>(.*?)</tr>", tabla, re.S)

        def celdas(tr):
            salida = []
            for attrs, contenido in re.findall(r"<t[hd]\b([^>]*)>(.*?)</t[hd]>", tr, re.S):
                m = re.search(r'colspan="(\d+)"', attrs)
                salida += [re.sub(r"<[^>]+>", "", contenido).strip()] + [""] * ((int(m.group(1)) if m else 1) - 1)
            return salida
        fechas, alb = celdas(filas[0]), celdas(filas[2])
        self.assertEqual(fechas[fechas.index("01/10")], "01/10")
        self.assertEqual(alb[fechas.index("01/10")], "111")
        self.assertEqual(alb[fechas.index("05/10")], "222")

    def test_el_mes_mas_reciente_va_primero(self):
        h = self._html()
        self.assertLess(h.index("Noviembre26"), h.index("Octubre26"))

    def test_sin_dato_se_pinta_con_guion_nunca_con_cero(self):
        h = self._html([_fila(datetime(2026, 10, 1), None, None, "T", "PVC", None, None, 20, 5)])
        self.assertIn("—", h)

    def test_escapa_el_texto_y_no_lleva_javascript(self):
        h = self._html([_fila(datetime(2026, 10, 1), "<b>Sumi</b>", "1", "T", "<script>x</script>", None, None, 1, 5)])
        self.assertNotIn("<script", h)
        self.assertNotIn("<b>Sumi</b>", h)
        self.assertIn("&lt;", h)

    def test_generar_panel_usa_la_vista_nueva(self):
        res = em.leer_materiales_wb(_wb(), hoy=HOY)
        carpeta = tempfile.mkdtemp(dir=None)
        try:
            salida = os.path.join(carpeta, "panel.html")
            panel_obra.generar_panel(
                obra="Obra", subtitulo="", historial=[], materiales=res,
                ficha={'_disponible': True, 'datos': {}, 'personal': [], 'hitos': [],
                       'riesgos': [], 'plan': [], 'tareas': []},
                documentos=[], prioridades={
                    'sin_base': False, 'revision': '28/07/2026', 'version': '4.3',
                    'catalogo_version': '1.3', 'resumen': {}, 'items': [], 'inventario': [],
                    'dudas_pendientes': [], 'preguntas_orden': [], 'prevision': [], 'avisos': []},
                output_path=salida)
            with open(salida, encoding="utf-8") as f:
                html = f.read()
        finally:
            import shutil
            shutil.rmtree(carpeta, ignore_errors=True)
        self.assertIn("Sumi B", html)
        self.assertIn("Albarán", html)

    def test_obras_con_el_formato_antiguo_siguen_igual(self):
        antiguo = {'disponible': True, 'meses': ['Mayo26'], 'ultimo_mes': 'Mayo26',
                   'ultima_fecha': '08/05/2026', 'dias_desde': 1, 'aviso': None,
                   'items': [{'categoria': 'CABLE', 'material': 'AZUL', 'tipo': '1.5', 'uni': '200', 'total': 9200}]}
        carpeta = tempfile.mkdtemp()
        try:
            salida = os.path.join(carpeta, "panel.html")
            panel_obra.generar_panel(
                obra="Obra", subtitulo="", historial=[], materiales=antiguo,
                ficha={'_disponible': True, 'datos': {}, 'personal': [], 'hitos': [],
                       'riesgos': [], 'plan': [], 'tareas': []},
                documentos=[], prioridades={
                    'sin_base': False, 'revision': '28/07/2026', 'version': '4.3',
                    'catalogo_version': '1.3', 'resumen': {}, 'items': [], 'inventario': [],
                    'dudas_pendientes': [], 'preguntas_orden': [], 'prevision': [], 'avisos': []},
                output_path=salida)
            with open(salida, encoding="utf-8") as f:
                html = f.read()
        finally:
            import shutil
            shutil.rmtree(carpeta, ignore_errors=True)
        self.assertIn("Consumo acumulado de material", html)
        self.assertIn("9.200", html)


class TestExcelGenerado(unittest.TestCase):

    def _libro(self, tmp, filas=None):
        ruta = os.path.join(tmp, "m.xlsx")
        ent, _ = em.leer_registro(_wb(filas))
        em.crear_libro(ent, ruta)
        return ruta

    def test_orden_de_hojas_y_nombres(self):
        with tempfile.TemporaryDirectory() as tmp:
            wb = _cargar(self._libro(tmp))
        self.assertEqual(wb.sheetnames, ["ENTREGAS", "Octubre26", "Noviembre26", "ACUMULADO", "LEEME"])

    def test_hoja_de_mes_lleva_fechas_suministrador_albaran_y_totales(self):
        with tempfile.TemporaryDirectory() as tmp:
            ws = _cargar(self._libro(tmp))["Octubre26"]
            g = _grid(ws)
        cab = g[0]
        self.assertEqual(cab[:5], ["MATERIAL", "TIPO", "MARCA", "UNI", "MES ANT"])
        self.assertEqual(cab[5:8], [datetime(2026, 10, 1), datetime(2026, 10, 5), datetime(2026, 10, 5)])
        self.assertEqual(cab[8:], ["TOTAL MES", "ACUMULADO"])
        self.assertEqual(g[1][0], "Suministrador")
        self.assertEqual(g[1][5:8], ["Sumi A", "Sumi B", "Sumi B"])
        self.assertEqual(g[2][0], "Albarán")
        self.assertEqual(g[2][5:8], ["111", "222", "223"])
        pvc = next(f for f in g if f[0] == "PVC")
        self.assertEqual(pvc[5:], [None, 171, 29, 200, 200])

    def test_las_hojas_generadas_no_llevan_formulas(self):
        with tempfile.TemporaryDirectory() as tmp:
            wb = _cargar(self._libro(tmp), valores=False)
        for nombre in ("Octubre26", "Noviembre26", "ACUMULADO"):
            for fila in wb[nombre].iter_rows():
                for c in fila:
                    self.assertFalse(isinstance(c.value, str) and c.value.startswith("="), (nombre, c.coordinate))

    def test_regenerar_es_idempotente_y_no_toca_entregas(self):
        with tempfile.TemporaryDirectory() as tmp:
            ruta = self._libro(tmp)
            antes = {n: _grid(_cargar(ruta)[n]) for n in ("ENTREGAS", "Octubre26", "ACUMULADO")}
            em.regenerar_hojas(ruta)
            em.regenerar_hojas(ruta)
            despues = {n: _grid(_cargar(ruta)[n]) for n in ("ENTREGAS", "Octubre26", "ACUMULADO")}
        self.assertEqual(antes, despues)

    def test_regenerar_recoge_las_filas_nuevas_y_respeta_hojas_ajenas(self):
        with tempfile.TemporaryDirectory() as tmp:
            ruta = self._libro(tmp)
            wb = openpyxl.load_workbook(ruta)
            wb.create_sheet("MIS NOTAS")["A1"] = "no tocar"
            wb["ENTREGAS"].append(_fila(datetime(2026, 12, 2), "Sumi C", "444", "TUBO", "PVC", "GRIS", None, 20, 70))
            wb.save(ruta)
            em.regenerar_hojas(ruta)
            wb2 = _cargar(ruta)
        self.assertIn("Diciembre26", wb2.sheetnames)
        self.assertEqual(wb2["MIS NOTAS"]["A1"].value, "no tocar")

    def test_detecta_hojas_desactualizadas(self):
        with tempfile.TemporaryDirectory() as tmp:
            ruta = self._libro(tmp)
            self.assertEqual(em.hojas_desactualizadas(_cargar(ruta)), [])
            wb = openpyxl.load_workbook(ruta)
            wb["ENTREGAS"].append(_fila(datetime(2026, 10, 9), "Sumi A", "9", "TUBO", "PVC", "GRIS", None, 20, 5))
            wb.save(ruta)
            self.assertTrue(em.hojas_desactualizadas(_cargar(ruta)))


class TestMigracionDesdeLaRejilla(unittest.TestCase):

    def _rejilla(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Octubre26"
        ws.append(["MATERIAL", "TIPO", "MARCA", "UNI", "MES ANT", datetime(2026, 10, 1), datetime(2026, 10, 5), "TOTAL"])
        ws.append([None] * 8)
        ws.append(["CABLE 1,5", None, None, None, None, None, None, None])
        ws.append(["AZUL", 1.5, None, 200, 0, 400, None, 400])
        ws.append(["GRIS", 1.5, None, 200, 0, None, 100, 100])
        ws.append(["NEGRO", 1.5, None, 200, 0, 0, None, 0])
        ws.append(["TUBO", None, None, None, None, None, None, None])
        ws.append(["PVC", "GRIS", None, 20, 7, 5, 6, 18])
        return wb

    def test_cada_celda_con_cantidad_es_una_entrega_y_mes_ant_no(self):
        ent = em.entregas_desde_rejilla(self._rejilla())
        claves = sorted((e["fecha"], e["categoria"], e["material"], e["cantidad"]) for e in ent)
        self.assertEqual(claves, [
            (D(2026, 10, 1), "CABLE 1,5", "AZUL", 400),
            (D(2026, 10, 1), "TUBO", "PVC", 5),
            (D(2026, 10, 5), "CABLE 1,5", "GRIS", 100),
            (D(2026, 10, 5), "TUBO", "PVC", 6),
        ])
        self.assertTrue(all(e["suministrador"] is None and e["albaran"] is None for e in ent))

    def test_las_entregas_migradas_salen_en_orden_cronologico_y_filas_consecutivas(self):
        """El registro lo va a leer y ampliar una persona: cronologico, y dentro
        de cada dia en el orden del catalogo de la rejilla."""
        ent = em.entregas_desde_rejilla(self._rejilla())
        self.assertEqual([e["fecha"] for e in ent], sorted(e["fecha"] for e in ent))
        self.assertEqual([e["fila"] for e in ent], list(range(2, 2 + len(ent))))
        del_primer_dia = [e["material"] for e in ent if e["fecha"] == D(2026, 10, 1)]
        self.assertEqual(del_primer_dia, ["AZUL", "PVC"])       # AZUL esta antes que PVC en la rejilla

    def test_migrar_conserva_el_total_del_excel_original(self):
        """La migracion no puede perder ni inventar cantidad: el total por
        material de la rejilla (sin MES ANT) tiene que ser el del registro."""
        ent = em.entregas_desde_rejilla(self._rejilla())
        g = em.agrupar(ent)
        total = {f["material"]: f["total"] for f in g["acumulado"]["filas"]}
        self.assertEqual(total, {"AZUL": 400, "GRIS": 100, "PVC": 11})


class TestHerramientaMaterialesXlsx(unittest.TestCase):

    def test_comprobar_un_libro_de_entregas_al_dia_esta_ok(self):
        with tempfile.TemporaryDirectory() as tmp:
            ruta = os.path.join(tmp, "m.xlsx")
            em.crear_libro(em.leer_registro(_wb(MUESTRA[:4]))[0], ruta)
            self.assertEqual(mx.diagnosticar(ruta), [])

    def test_comprobar_avisa_si_las_hojas_no_estan_al_dia(self):
        with tempfile.TemporaryDirectory() as tmp:
            ruta = os.path.join(tmp, "m.xlsx")
            em.crear_libro(em.leer_registro(_wb(MUESTRA[:4]))[0], ruta)
            wb = openpyxl.load_workbook(ruta)
            wb["ENTREGAS"].append(_fila(datetime(2026, 10, 9), "S", "9", "TUBO", "PVC", "GRIS", None, 20, 5))
            wb.save(ruta)
            defectos = mx.diagnosticar(ruta)
        self.assertTrue(any("generar" in d for d in defectos), defectos)

    def test_comprobar_cuenta_las_filas_invalidas_como_defecto(self):
        with tempfile.TemporaryDirectory() as tmp:
            ruta = os.path.join(tmp, "m.xlsx")
            wb = _wb(MUESTRA[:1] + [_fila("ayer", "S", "1", "T", "PVC", None, None, 20, 5)])
            wb.save(ruta)
            defectos = mx.diagnosticar(ruta)
        self.assertTrue(any("fila 3" in d for d in defectos), defectos)

    def test_generar_deja_el_libro_al_dia_y_devuelve_cero(self):
        with tempfile.TemporaryDirectory() as tmp:
            ruta = os.path.join(tmp, "REVISIONES", "m.xlsx")
            os.makedirs(os.path.dirname(ruta))
            em.crear_libro(em.leer_registro(_wb(MUESTRA[:4]))[0], ruta)
            wb = openpyxl.load_workbook(ruta)
            wb["ENTREGAS"].append(_fila(datetime(2026, 10, 9), "S", "9", "TUBO", "PVC", "GRIS", None, 20, 5))
            wb.save(ruta)
            salida = io.StringIO()
            from contextlib import redirect_stdout
            with redirect_stdout(salida):
                codigo = mx.main(["generar", ruta])
            defectos = mx.diagnosticar(ruta)
            copias = os.listdir(os.path.join(tmp, "REVISIONES", "_SISTEMA"))
        self.assertEqual(codigo, 0, salida.getvalue())
        self.assertEqual(defectos, [])
        self.assertEqual(len(copias), 1)           # copia de seguridad previa

    def test_migrar_convierte_la_rejilla_en_registro_con_copia(self):
        with tempfile.TemporaryDirectory() as tmp:
            ruta = os.path.join(tmp, "REVISIONES", "m.xlsx")
            os.makedirs(os.path.dirname(ruta))
            TestMigracionDesdeLaRejilla()._rejilla().save(ruta)
            from contextlib import redirect_stdout
            with redirect_stdout(io.StringIO()):
                codigo = mx.main(["migrar", ruta])
            wb = _cargar(ruta)
            copias = os.listdir(os.path.join(tmp, "REVISIONES", "_SISTEMA"))
            defectos = mx.diagnosticar(ruta)
        self.assertEqual(codigo, 0)
        self.assertIn("ENTREGAS", wb.sheetnames)
        self.assertEqual(wb["ENTREGAS"].max_row, 5)       # cabecera + 4 entregas
        self.assertEqual(defectos, [])
        self.assertEqual(len(copias), 1)

    def test_migrar_aborta_sin_tocar_nada_si_se_perderia_una_cantidad(self):
        """Una fila con cantidad pero sin nombre no se puede convertir en
        entrega: la migracion tiene que parar en vez de dar por buena una
        rejilla a la que le faltan unidades."""
        wb = TestMigracionDesdeLaRejilla()._rejilla()
        wb.active.append([None, None, None, None, 0, 999, None, 999])      # cantidad sin material
        with tempfile.TemporaryDirectory() as tmp:
            ruta = os.path.join(tmp, "REVISIONES", "m.xlsx")
            os.makedirs(os.path.dirname(ruta))
            wb.save(ruta)
            antes = open(ruta, "rb").read()
            from contextlib import redirect_stdout
            with redirect_stdout(io.StringIO()):
                codigo = mx.main(["migrar", ruta])
            despues = open(ruta, "rb").read()
            hay_copia = os.path.exists(os.path.join(tmp, "REVISIONES", "_SISTEMA"))
        self.assertEqual(codigo, 1)
        self.assertEqual(antes, despues)
        self.assertFalse(hay_copia)

    def test_migrar_no_pisa_un_libro_que_ya_es_registro(self):
        with tempfile.TemporaryDirectory() as tmp:
            ruta = os.path.join(tmp, "m.xlsx")
            em.crear_libro(em.leer_registro(_wb())[0], ruta)
            from contextlib import redirect_stdout
            with redirect_stdout(io.StringIO()):
                self.assertEqual(mx.main(["migrar", ruta]), 1)


if __name__ == "__main__":
    unittest.main()
